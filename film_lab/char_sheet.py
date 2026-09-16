"""Character Bible sheets as local Word / PDF.

Bios are prose. Excel and PowerPoint stay on Writing Studio
(shot lists / decks) — they are awkward for a character sheet.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from film_lab.filming import CAST_ROLES, ROLE_ADULT, normalize_role
from film_lab.office import OfficeError, extract_office, file_paths
from film_lab.pdf_export import write_simple_pdf
from film_lab.util import slugify

BIBLE_IMPORT_SUFFIXES = {".docx", ".pdf"}
BIBLE_EXPORT_LABELS = ("Word", "PDF")
BIBLE_REJECT = {".xlsx", ".xls", ".pptx", ".ppt"}

SHEET_FIELDS: tuple[tuple[str, str], ...] = (
    ("id", "Id"),
    ("name", "Name"),
    ("role", "Role"),
    ("age_band", "Age band"),
    ("age_years", "Age (years)"),
    ("look_notes", "Look / appearance"),
    ("wardrobe", "Wardrobe"),
    ("rings_props", "Rings / props"),
    ("personality", "Personality"),
    ("emotion_baseline", "Emotion baseline"),
    ("micro_expression", "Micro-expression"),
    ("behavior", "Behavior"),
    ("lighting_notes", "Lighting"),
    ("voice_notes", "Voice notes"),
    ("locked_descriptor", "Locked descriptor"),
    ("voice_backend", "Voice backend"),
    ("voice_id", "Voice id"),
)

_ALIASES: dict[str, str] = {
    "id": "id",
    "character id": "id",
    "name": "name",
    "character": "name",
    "character name": "name",
    "role": "role",
    "role tag": "role",
    "age band": "age_band",
    "age": "age_band",
    "age (years)": "age_years",
    "age years": "age_years",
    "years": "age_years",
    "look": "look_notes",
    "look / appearance": "look_notes",
    "appearance": "look_notes",
    "bio": "look_notes",
    "biography": "look_notes",
    "description": "look_notes",
    "wardrobe": "wardrobe",
    "costume": "wardrobe",
    "rings / props": "rings_props",
    "rings": "rings_props",
    "props": "rings_props",
    "personality": "personality",
    "emotion": "emotion_baseline",
    "emotion baseline": "emotion_baseline",
    "micro-expression": "micro_expression",
    "micro expression": "micro_expression",
    "behavior": "behavior",
    "lighting": "lighting_notes",
    "lighting notes": "lighting_notes",
    "voice": "voice_notes",
    "voice notes": "voice_notes",
    "voice / performance": "voice_notes",
    "locked descriptor": "locked_descriptor",
    "descriptor": "locked_descriptor",
    "voice backend": "voice_backend",
    "tts backend": "voice_backend",
    "voice id": "voice_id",
    "tts voice id": "voice_id",
}

_LABELED = re.compile(
    r"^(?:[-*]\s+)?(?:\*\*)?([^:*#]+?)(?:\*\*)?\s*:\s+(.*)$"
)
_HEADING = re.compile(r"^#{1,3}\s+(.+)$")
_TITLE = re.compile(
    r"^(?:character\s+sheet|character\s+bible|bio)\s+[—–-]\s+(.+)$",
    re.I,
)
_SCALAR = {
    "id",
    "name",
    "role",
    "age_band",
    "age_years",
    "micro_expression",
    "behavior",
    "voice_backend",
    "voice_id",
}


class BibleSheetError(RuntimeError):
    """Could not read or write a Character Bible sheet."""


def ext_for_bible_export(raw: str | None) -> str:
    key = (raw or "Word").strip().lower()
    if key in {"word", "docx", ".docx"}:
        return ".docx"
    if key in {"pdf", ".pdf"}:
        return ".pdf"
    raise BibleSheetError("Character Bible export is Word or PDF only.")


def format_sheet(fields: dict[str, Any], *, title: str | None = None) -> str:
    name = str(fields.get("name") or "Character").strip() or "Character"
    heading = (title or f"Character sheet — {name}").strip()
    lines = [
        heading,
        "",
        "Film Lab Character Bible. Local file. Zero credits. Not a spreadsheet.",
        "",
    ]
    for key, label in SHEET_FIELDS:
        value = fields.get(key)
        if value is None:
            continue
        text = str(value).strip()
        if not text:
            continue
        lines.append(f"{label}: {text}")
    return "\n".join(lines).strip() + "\n"


def parse_sheet_text(text: str) -> dict[str, str]:
    """Read a labeled bio. Unlabeled leftover becomes look / appearance."""
    body = (text or "").replace("\r\n", "\n").strip()
    if not body:
        raise BibleSheetError("That file has no extractable text.")
    found: dict[str, list[str]] = {}
    current: str | None = None
    leftover: list[str] = []

    for raw in body.splitlines():
        line = raw.strip()
        if not line or line in {"---", "***"}:
            continue
        if line.lower().startswith("film lab character bible"):
            continue
        titled = _TITLE.match(_strip_hashes(line))
        if titled and "name" not in found:
            _push_field(found, "name", titled.group(1).strip())
            current = None
            continue
        key, value = _match_label(line)
        if key:
            if value:
                _push_field(found, key, value)
                current = None if key in _SCALAR else key
            else:
                current = key
            continue
        if current:
            _push_field(found, current, line)
            if current in _SCALAR:
                current = None
        else:
            leftover.append(line)

    out: dict[str, str] = {}
    for key, parts in found.items():
        joined = "\n".join(p for p in parts if p).strip()
        if joined:
            out[key] = joined
    if leftover and "look_notes" not in out:
        # First short leftover line is often the name.
        if "name" not in out and leftover[0] and len(leftover[0]) <= 48:
            out["name"] = leftover[0]
            leftover = leftover[1:]
        if leftover:
            out["look_notes"] = "\n".join(leftover).strip()
    if "role" in out:
        out["role"] = normalize_role(out["role"])
        if out["role"] not in CAST_ROLES:
            out["role"] = ROLE_ADULT
    if "age_years" in out:
        digits = re.search(r"\d+", out["age_years"])
        if digits:
            out["age_years"] = digits.group(0)
    if not out:
        raise BibleSheetError("Could not read a name or bio from that file.")
    return out


def import_bible_files(files) -> tuple[dict[str, str], list[str]]:
    paths = file_paths(files)
    if not paths:
        raise BibleSheetError("Upload a Word (.docx) or PDF character sheet first.")
    merged: dict[str, str] = {}
    names: list[str] = []
    notes: list[str] = []
    for path in paths:
        try:
            parsed = parse_sheet_text(_extract_bible(path))
        except (BibleSheetError, OfficeError) as exc:
            notes.append(str(exc))
            continue
        names.append(path.name)
        merged.update(parsed)
    if not merged:
        raise BibleSheetError(
            notes[0] if notes else "Those files had no character sheet text."
        )
    return merged, names


def export_bible_sheet(
    fields: dict[str, Any],
    dest: Path,
    *,
    title: str | None = None,
) -> Path:
    dest = Path(dest)
    suffix = dest.suffix.lower()
    body = format_sheet(fields, title=title)
    name = str(fields.get("name") or "Character").strip() or "Character"
    heading = title or f"Character sheet — {name}"
    if suffix == ".docx":
        return _write_bible_docx(body, dest, title=heading)
    if suffix == ".pdf":
        dest.parent.mkdir(parents=True, exist_ok=True)
        return write_simple_pdf(body, dest, title=heading)
    raise BibleSheetError("Character Bible export is Word (.docx) or PDF only.")


def sheet_filename(name: str, cid: str, ext: str) -> str:
    slug = slugify(name or cid or "character", cid or "character")
    return f"{slug}-character-sheet{ext}"


def _extract_bible(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in BIBLE_REJECT:
        raise BibleSheetError(
            "Character Bible is Word or PDF only. "
            "Excel / PowerPoint stay on Writing Studio (shot lists / decks)."
        )
    if suffix not in BIBLE_IMPORT_SUFFIXES:
        raise BibleSheetError(
            f"Cannot import `{path.name}` into Character Bible. Use Word (.docx) or PDF."
        )
    return extract_office(path)


def _match_label(line: str) -> tuple[str | None, str]:
    heading = _HEADING.match(line)
    if heading:
        key = _ALIASES.get(heading.group(1).strip().lower().rstrip(":"))
        return key, ""
    labeled = _LABELED.match(line)
    if not labeled:
        return None, ""
    key = _ALIASES.get(labeled.group(1).strip().lower())
    if not key:
        return None, ""
    return key, labeled.group(2).strip()


def _push_field(found: dict[str, list[str]], key: str, value: str) -> None:
    text = (value or "").strip()
    if not text:
        return
    if key in _SCALAR:
        found[key] = [text]
        return
    found.setdefault(key, []).append(text)


def _strip_hashes(line: str) -> str:
    return re.sub(r"^#+\s*", "", line).strip()


def _write_bible_docx(text: str, dest: Path, *, title: str) -> Path:
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError as exc:
        raise BibleSheetError(
            "Word export needs python-docx. pip install python-docx  "
            "(already in requirements.txt)."
        ) from exc
    doc = Document()
    heading = doc.add_heading(title, level=0)
    for run in heading.runs:
        run.font.size = Pt(22)
    for raw in text.splitlines():
        line = raw.strip()
        if not line or line == title:
            continue
        labeled = _LABELED.match(line)
        if labeled:
            para = doc.add_paragraph()
            run = para.add_run(f"{labeled.group(1).strip()}: ")
            run.bold = True
            para.add_run(labeled.group(2).strip())
        else:
            doc.add_paragraph(line)
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest))
    return dest
