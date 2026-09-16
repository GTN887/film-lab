"""Local office files for Writing Studio.

Import / view: PDF, Word (.docx), PowerPoint (.pptx), Excel (.xlsx),
plus .txt / .md. Export what they type: Word, PDF, PowerPoint, Excel,
plain text. All on this machine. No cloud convert.

Missing library → toast, not a crash. Scanned PDFs have no OCR yet.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from film_lab.pdf_export import write_simple_pdf

ALLOWED_IMPORT_SUFFIXES = {
    ".txt",
    ".md",
    ".markdown",
    ".pdf",
    ".docx",
    ".pptx",
    ".xlsx",
}
RTF_LATER = {".rtf"}

EXPORT_LABELS = ("Word", "PDF", "PowerPoint", "Excel", "plain text")
EXPORT_LABEL_TO_EXT = {
    "word": ".docx",
    "docx": ".docx",
    ".docx": ".docx",
    "pdf": ".pdf",
    ".pdf": ".pdf",
    "powerpoint": ".pptx",
    "pptx": ".pptx",
    ".pptx": ".pptx",
    "excel": ".xlsx",
    "xlsx": ".xlsx",
    ".xlsx": ".xlsx",
    "plain text": ".txt",
    "text": ".txt",
    "txt": ".txt",
    ".txt": ".txt",
    "fountain": ".fountain",
    ".fountain": ".fountain",
    "md": ".md",
    "markdown": ".md",
    ".md": ".md",
}

_SLUG_LINE = re.compile(r"^(INT\.|EXT\.|INT/EXT\.|I/E\.)", re.I)
_BEAT_MARK = re.compile(r"^\*\*BEAT\s+\d+", re.I)
_NUMBERED = re.compile(r"^(?:BEAT\s+)?\d+[\.\):]\s+\S", re.I)
_HEADING = re.compile(r"^#{1,3}\s+\S")
_SLIDE_CAP = 1600


class OfficeError(RuntimeError):
    """Could not read or write an office file. Toast this — do not crash."""


def ext_for_export_label(raw: str | None) -> str:
    key = (raw or "Word").strip()
    mapped = EXPORT_LABEL_TO_EXT.get(key) or EXPORT_LABEL_TO_EXT.get(key.lower())
    if mapped:
        return mapped
    raise OfficeError(
        "Export as Word, PDF, PowerPoint, Excel, or plain text "
        "(fountain / md still work)."
    )


def extract_office(path: str | Path) -> str:
    """Read one upload into viewable text."""
    src = Path(path)
    if not src.is_file():
        raise OfficeError(f"File not found: {src.name or src}")
    suffix = src.suffix.lower()
    if suffix in RTF_LATER:
        raise OfficeError(
            "RTF later — use PDF, Word (.docx), PowerPoint (.pptx), "
            "Excel (.xlsx), or text."
        )
    if suffix not in ALLOWED_IMPORT_SUFFIXES:
        raise OfficeError(
            f"Cannot import `{src.name}`. Use PDF, Word (.docx), "
            "PowerPoint (.pptx), Excel (.xlsx), or .txt / .md."
        )
    if suffix in {".txt", ".md", ".markdown"}:
        return _read_plain(src)
    if suffix == ".pdf":
        return _read_pdf(src)
    if suffix == ".docx":
        return _read_docx(src)
    if suffix == ".pptx":
        return _read_pptx(src)
    return _read_xlsx(src)


def import_files_as_text(files) -> tuple[str, list[str]]:
    """Upload → one draft body (view). Multiple files join with headers."""
    paths = file_paths(files)
    if not paths:
        raise OfficeError(
            "Upload a PDF, Word, PowerPoint, Excel, or text file first."
        )
    chunks: list[str] = []
    names: list[str] = []
    notes: list[str] = []
    for path in paths:
        try:
            text = extract_office(path).strip()
        except OfficeError as exc:
            notes.append(str(exc))
            continue
        if text:
            names.append(path.name)
            if len(paths) > 1:
                chunks.append(f"# {path.name}\n\n{text}")
            else:
                chunks.append(text)
        else:
            notes.append(f"`{path.name}` had no extractable text.")
    if not chunks:
        raise OfficeError(notes[0] if notes else "Those files had no extractable text.")
    body = "\n\n---\n\n".join(chunks)
    if notes:
        body += "\n\n# Import notes\n\n" + "\n".join(f"- {n}" for n in notes)
    return body, names


def export_typed(
    text: str,
    dest: Path,
    *,
    title: str = "Draft",
) -> Path:
    """Write the typed draft to Word, PDF, PowerPoint, Excel, or plain text."""
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    suffix = dest.suffix.lower()
    body = text if text is not None else ""
    name = (title or "Draft").strip() or "Draft"
    if suffix == ".docx":
        return write_word(body, dest, title=name)
    if suffix == ".pdf":
        return write_pdf(body, dest, title=name)
    if suffix == ".pptx":
        return write_pptx(body, dest, title=name)
    if suffix == ".xlsx":
        return write_xlsx(body, dest, title=name)
    if suffix == ".txt":
        dest.write_text(body, encoding="utf-8")
        return dest
    raise OfficeError(
        "Export must be Word (.docx), PDF, PowerPoint (.pptx), "
        "Excel (.xlsx), or plain text (.txt)."
    )


def write_word(text: str, dest: Path, *, title: str = "Draft") -> Path:
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError as exc:
        raise OfficeError(
            "Word export needs python-docx. pip install python-docx  "
            "(already in requirements.txt)."
        ) from exc
    doc = Document()
    heading = doc.add_heading(title, level=0)
    for run in heading.runs:
        run.font.size = Pt(22)
    pages = pages_from_text(text)
    if not (text or "").strip():
        doc.add_paragraph("(empty draft)")
    else:
        for heading_text, body in pages:
            if heading_text:
                doc.add_heading(heading_text, level=1)
            for para in _paragraphs(body):
                doc.add_paragraph(para)
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dest))
    return dest


def write_pdf(text: str, dest: Path, *, title: str = "Draft") -> Path:
    pages = pages_from_text(text)
    if (text or "").strip():
        blocks = [f"{title}\n"]
        for heading, body in pages:
            if heading:
                blocks.append(heading)
            if body:
                blocks.append(body)
        blob = "\n\n".join(blocks).strip() + "\n"
    else:
        blob = f"{title}\n\n(empty draft)\n"
    return write_simple_pdf(blob, dest, title=title)


def write_pptx(text: str, dest: Path, *, title: str = "Draft") -> Path:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise OfficeError(
            "PowerPoint export needs python-pptx. pip install python-pptx  "
            "(already in requirements.txt)."
        ) from exc
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    title_layout = prs.slide_layouts[0]
    body_layout = prs.slide_layouts[1]
    opener = prs.slides.add_slide(title_layout)
    _set_shape_text(opener.shapes.title, title)
    if len(opener.placeholders) > 1:
        _set_shape_text(
            opener.placeholders[1],
            "Film Lab Writing Studio — local export. Zero credits.",
        )
    pages = pages_from_text(text)
    if not (text or "").strip():
        slide = prs.slides.add_slide(body_layout)
        _set_shape_text(slide.shapes.title, "Draft")
        if len(slide.placeholders) > 1:
            _set_shape_text(slide.placeholders[1], "(empty draft)")
    else:
        for heading, body in pages:
            chunks = _chunk_slide(body)
            for index, chunk in enumerate(chunks):
                slide = prs.slides.add_slide(body_layout)
                label = heading if index == 0 else f"{heading} (cont.)"
                _set_shape_text(slide.shapes.title, label[:120] or f"Beat {index + 1}")
                if len(slide.placeholders) > 1:
                    tf = slide.placeholders[1].text_frame
                    tf.clear()
                    tf.word_wrap = True
                    p = tf.paragraphs[0]
                    run = p.add_run()
                    run.text = chunk
                    run.font.size = Pt(18)
    dest.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(dest))
    return dest


def write_xlsx(text: str, dest: Path, *, title: str = "Draft") -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise OfficeError(
            "Excel export needs openpyxl. pip install openpyxl  "
            "(already in requirements.txt)."
        ) from exc
    wb = Workbook()
    ws = wb.active
    sheet_name = re.sub(r"[\\/*?:\[\]]", " ", title)[:31].strip() or "Draft"
    ws.title = sheet_name
    ws.append(["Beat", "Heading", "Body"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    pages = pages_from_text(text)
    if not (text or "").strip():
        ws.append([1, title, "(empty draft)"])
    else:
        for index, (heading, body) in enumerate(pages, start=1):
            ws.append([index, heading, body])
    widths = (8, 36, 80)
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(dest))
    return dest


def pages_from_text(text: str) -> list[tuple[str, str]]:
    """Split typed pages into (heading, body) for slides and spreadsheet rows."""
    body = (text or "").strip()
    if not body:
        return [("Draft", "")]
    lines = body.splitlines()
    blocks: list[str]
    if any(_SLUG_LINE.match(line) for line in lines):
        blocks = _split_on(_SLUG_LINE, lines)
    elif sum(1 for line in lines if _BEAT_MARK.match(line)) >= 2:
        blocks = _split_on(_BEAT_MARK, lines)
    elif sum(1 for line in lines if _NUMBERED.match(line)) >= 2:
        blocks = _split_on(_NUMBERED, lines)
    elif sum(1 for line in lines if _HEADING.match(line)) >= 2:
        blocks = _split_on(_HEADING, lines)
    else:
        blocks = [p.strip() for p in re.split(r"\n\s*\n", body) if p.strip()] or [body]
    pages: list[tuple[str, str]] = []
    for index, block in enumerate(blocks, start=1):
        first, _, rest = block.strip().partition("\n")
        first = first.strip()
        rest = rest.strip()
        if (
            _SLUG_LINE.match(first)
            or _BEAT_MARK.match(first)
            or _HEADING.match(first)
            or _NUMBERED.match(first)
        ):
            heading = first.lstrip("#").strip()
            pages.append((heading or f"Beat {index}", rest))
        else:
            pages.append((f"Beat {index}", block.strip()))
    return pages or [("Draft", body)]


def file_paths(files) -> list[Path]:
    if not files:
        return []
    if not isinstance(files, list):
        files = [files]
    out: list[Path] = []
    for item in files:
        if item is None:
            continue
        if isinstance(item, dict):
            raw = item.get("path") or item.get("name") or ""
            if not raw:
                continue
            path = Path(raw)
        elif isinstance(item, (str, Path)):
            path = Path(item)
        elif hasattr(item, "name"):
            path = Path(item.name)
        else:
            continue
        if path.is_file():
            out.append(path)
    return out


def _read_plain(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise OfficeError(
            "PDF import needs pypdf. pip install pypdf  (already in requirements.txt)."
        ) from exc
    try:
        reader = PdfReader(str(path))
    except Exception as exc:  # noqa: BLE001 — surface bad PDF honestly
        raise OfficeError(f"Could not read PDF `{path.name}`: {exc}") from exc
    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    text = "\n\n".join(p.strip() for p in pages if p and p.strip())
    if not text:
        raise OfficeError(
            f"`{path.name}` has no extractable text (scanned PDF needs OCR later)."
        )
    return text


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise OfficeError(
            "Word import needs python-docx. pip install python-docx  "
            "(already in requirements.txt)."
        ) from exc
    try:
        doc = Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise OfficeError(f"Could not read Word `{path.name}`: {exc}") from exc
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text and para.text.strip():
            parts.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n\n".join(parts)
    if not text:
        raise OfficeError(f"`{path.name}` has no extractable text.")
    return text


def _read_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise OfficeError(
            "PowerPoint import needs python-pptx. pip install python-pptx  "
            "(already in requirements.txt)."
        ) from exc
    try:
        prs = Presentation(str(path))
    except Exception as exc:  # noqa: BLE001
        raise OfficeError(f"Could not read PowerPoint `{path.name}`: {exc}") from exc
    slides: list[str] = []
    for index, slide in enumerate(prs.slides, start=1):
        seen: set[str] = set()
        parts: list[str] = []
        title = ""
        if slide.shapes.title is not None:
            title = (slide.shapes.title.text or "").strip()
            if title:
                parts.append(title)
                seen.add(title)
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                table = shape.table
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            if not getattr(shape, "has_text_frame", False):
                continue
            blob = (shape.text_frame.text or "").strip()
            if blob and blob not in seen:
                seen.add(blob)
                if blob != title:
                    parts.append(blob)
        if parts:
            slides.append(f"## Slide {index}\n\n" + "\n\n".join(parts))
    text = "\n\n".join(slides).strip()
    if not text:
        raise OfficeError(f"`{path.name}` has no extractable text.")
    return text


def _read_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise OfficeError(
            "Excel import needs openpyxl. pip install openpyxl  "
            "(already in requirements.txt)."
        ) from exc
    try:
        wb = load_workbook(str(path), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise OfficeError(f"Could not read Excel `{path.name}`: {exc}") from exc
    sheets: list[str] = []
    try:
        for ws in wb.worksheets:
            rows: list[str] = []
            for row in ws.iter_rows(values_only=True):
                cells = [_cell_text(cell) for cell in row]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                sheets.append(f"## {ws.title}\n\n" + "\n".join(rows))
    finally:
        wb.close()
    text = "\n\n".join(sheets).strip()
    if not text:
        raise OfficeError(f"`{path.name}` has no extractable text.")
    return text


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _paragraphs(text: str) -> Iterable[str]:
    body = (text or "").replace("\r\n", "\n")
    paras = [p.strip() for p in re.split(r"\n\s*\n", body) if p.strip()]
    if paras:
        return paras
    lines = [ln.rstrip() for ln in body.split("\n") if ln.strip()]
    return lines or [""]


def _split_on(pattern: re.Pattern[str], lines: list[str]) -> list[str]:
    chunks: list[list[str]] = []
    current: list[str] = []
    for line in lines:
        if pattern.match(line) and current:
            chunks.append(current)
            current = [line]
        else:
            current.append(line)
    if current:
        chunks.append(current)
    return ["\n".join(block).strip() for block in chunks if "\n".join(block).strip()]


def _chunk_slide(text: str) -> list[str]:
    body = (text or "").strip() or "(empty)"
    if len(body) <= _SLIDE_CAP:
        return [body]
    chunks: list[str] = []
    rest = body
    while rest:
        if len(rest) <= _SLIDE_CAP:
            chunks.append(rest)
            break
        cut = rest.rfind("\n", 0, _SLIDE_CAP)
        if cut < 80:
            cut = rest.rfind(" ", 0, _SLIDE_CAP)
        if cut < 80:
            cut = _SLIDE_CAP
        chunks.append(rest[:cut].strip())
        rest = rest[cut:].strip()
    return chunks or [body]


def _set_shape_text(shape, text: str) -> None:
    if shape is None:
        return
    try:
        shape.text = text
    except Exception:
        if getattr(shape, "text_frame", None) is not None:
            shape.text_frame.text = text
